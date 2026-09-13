#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Script xuất bản toàn diện 18 bài viết Bến Thành từ Markdown sang Microsoft Word (.docx)
Bao gồm:
- 18 bài Tiếng Việt riêng lẻ
- 18 bài Tiếng Anh riêng lẻ
- 1 Master Dossier Tổng Hợp 18 Bài Tiếng Việt
- 1 Master Compendium Tổng Hợp 18 Bài Tiếng Anh
Xuất bản đồng thời vào content-pipeline/campaign-ben-thanh/word/ và content-pipeline/06-word-documents/
"""

import os
import re
import markdown
from bs4 import BeautifulSoup, NavigableString, Tag
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

NAVY_HEX = "1B365D"
GOLD_HEX = "D97706"
SLATE_HEX = "475569"
LIGHT_BG_HEX = "F8FAFC"
BORDER_HEX = "CBD5E1"

COLOR_NAVY = RGBColor(27, 54, 93)      # #1B365D
COLOR_SLATE = RGBColor(71, 85, 105)    # #475569
COLOR_DARK = RGBColor(30, 41, 59)      # #1E293B
COLOR_AMBER = RGBColor(217, 119, 6)    # #D97706

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=120, bottom=120, left=180, right=180):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def set_table_borders(table):
    tblPr = table._element.xpath('w:tblPr')
    if tblPr:
        borders = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            f'<w:top w:val="single" w:sz="4" w:space="0" w:color="{BORDER_HEX}"/>'
            f'<w:bottom w:val="single" w:sz="6" w:space="0" w:color="{NAVY_HEX}"/>'
            f'<w:left w:val="none"/>'
            f'<w:right w:val="none"/>'
            f'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="{BORDER_HEX}"/>'
            f'<w:insideV w:val="none"/>'
            f'</w:tblBorders>'
        )
        tblPr[0].append(borders)

def parse_markdown_file(file_path):
    with open(file_path, "r", encoding="utf-8") as f:
        text = f.read()

    frontmatter = {}
    body = text
    if text.startswith("---"):
        parts = text.split("---", 2)
        if len(parts) >= 3:
            fm_text = parts[1]
            body = parts[2].strip()
            for line in fm_text.strip().split("\n"):
                if ":" in line:
                    key, val = line.split(":", 1)
                    key = key.strip()
                    val = val.strip().strip('"').strip("'")
                    frontmatter[key] = val

    return frontmatter, body

def style_paragraph(p, font_name="Calibri", font_size=11, color=COLOR_DARK, space_before=0, space_after=6, line_spacing=1.15):
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing
    for r in p.runs:
        r.font.name = font_name
        r.font.size = Pt(font_size)
        r.font.color.rgb = color

def add_callout_box(doc, text_content, label="CHÚ Ý / TRAVEL ADVISORY", is_gold=False):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    cell = table.cell(0, 0)
    cell.width = Inches(6.5)
    border_color = GOLD_HEX if is_gold else NAVY_HEX
    bg_color = "FFFBEB" if is_gold else "F1F5F9"
    set_cell_background(cell, bg_color)
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)

    tcPr = cell._element.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="none"/>'
        f'<w:bottom w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'<w:left w:val="single" w:sz="24" w:space="0" w:color="{border_color}"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)

    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    
    lbl_run = p.add_run(f"📌 {label}\n")
    lbl_run.font.name = "Calibri"
    lbl_run.font.size = Pt(9.5)
    lbl_run.font.bold = True
    lbl_run.font.color.rgb = COLOR_AMBER if is_gold else COLOR_NAVY

    txt_run = p.add_run(text_content.strip())
    txt_run.font.name = "Calibri"
    txt_run.font.size = Pt(10)
    txt_run.font.color.rgb = COLOR_DARK

    post_p = doc.add_paragraph()
    post_p.paragraph_format.space_before = Pt(0)
    post_p.paragraph_format.space_after = Pt(4)

def render_table_element(doc, table_soup):
    rows = table_soup.find_all("tr")
    if not rows:
        return

    num_cols = max(len(r.find_all(["th", "td"])) for r in rows)
    num_rows = len(rows)

    docx_table = doc.add_table(rows=num_rows, cols=num_cols)
    docx_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(docx_table)

    for r_idx, row in enumerate(rows):
        cells = row.find_all(["th", "td"])
        is_header = bool(row.find("th")) or (r_idx == 0)

        for c_idx, cell_data in enumerate(cells):
            if c_idx >= num_cols:
                break
            cell = docx_table.cell(r_idx, c_idx)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

            if is_header:
                set_cell_background(cell, NAVY_HEX)
                set_cell_margins(cell, top=140, bottom=140, left=150, right=150)
            else:
                bg = LIGHT_BG_HEX if r_idx % 2 == 1 else "FFFFFF"
                set_cell_background(cell, bg)
                set_cell_margins(cell, top=100, bottom=100, left=150, right=150)

            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            
            text = cell_data.get_text().strip()
            run = p.add_run(text)
            run.font.name = "Calibri"

            if is_header:
                run.font.bold = True
                run.font.size = Pt(9.5)
                run.font.color.rgb = RGBColor(255, 255, 255)
            else:
                run.font.size = Pt(9.5)
                run.font.color.rgb = COLOR_DARK

    post_p = doc.add_paragraph()
    post_p.paragraph_format.space_before = Pt(0)
    post_p.paragraph_format.space_after = Pt(4)

def parse_html_to_docx(doc, html_content):
    soup = BeautifulSoup(html_content, "html.parser")

    for elem in soup.children:
        if isinstance(elem, NavigableString):
            text = elem.strip()
            if text:
                p = doc.add_paragraph()
                p.add_run(text)
                style_paragraph(p)
            continue

        if not isinstance(elem, Tag):
            continue

        tag_name = elem.name.lower()

        if tag_name == "h1":
            continue

        elif tag_name == "h2":
            p = doc.add_paragraph()
            p.paragraph_format.keep_with_next = True
            run = p.add_run(elem.get_text().strip())
            run.font.name = "Calibri"
            run.font.size = Pt(16)
            run.font.bold = True
            run.font.color.rgb = COLOR_NAVY
            style_paragraph(p, font_name="Calibri", font_size=16, color=COLOR_NAVY, space_before=16, space_after=6)

        elif tag_name == "h3":
            p = doc.add_paragraph()
            p.paragraph_format.keep_with_next = True
            run = p.add_run(elem.get_text().strip())
            run.font.name = "Calibri"
            run.font.size = Pt(13)
            run.font.bold = True
            run.font.color.rgb = COLOR_AMBER
            style_paragraph(p, font_name="Calibri", font_size=13, color=COLOR_AMBER, space_before=12, space_after=4)

        elif tag_name == "h4":
            p = doc.add_paragraph()
            p.paragraph_format.keep_with_next = True
            run = p.add_run(elem.get_text().strip())
            run.font.name = "Calibri"
            run.font.size = Pt(11.5)
            run.font.bold = True
            run.font.color.rgb = COLOR_NAVY
            style_paragraph(p, font_name="Calibri", font_size=11.5, color=COLOR_NAVY, space_before=8, space_after=3)

        elif tag_name == "p":
            text = elem.get_text().strip()
            if not text:
                continue

            if text.startswith(">") or text.startswith("🏷️") or text.startswith("📍") or text.startswith("⚠️"):
                add_callout_box(doc, text, label="THÔNG TIN ĐẶC BIỆT")
                continue

            p = doc.add_paragraph()
            for child in elem.children:
                if isinstance(child, NavigableString):
                    p.add_run(str(child))
                elif isinstance(child, Tag):
                    r = p.add_run(child.get_text())
                    if child.name in ["strong", "b"]:
                        r.font.bold = True
                    if child.name in ["em", "i"]:
                        r.font.italic = True
                    if child.name == "a":
                        r.font.underline = True
                        r.font.color.rgb = COLOR_AMBER
            style_paragraph(p, font_name="Calibri", font_size=10.5, color=COLOR_DARK, space_before=2, space_after=6)

        elif tag_name == "blockquote":
            text = elem.get_text().strip()
            add_callout_box(doc, text, label="THE RICE TOUR ADVISORY", is_gold=True)

        elif tag_name in ["ul", "ol"]:
            for li in elem.find_all("li", recursive=False):
                p = doc.add_paragraph(style="List Bullet" if tag_name == "ul" else "List Number")
                p.paragraph_format.space_before = Pt(1)
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.15
                
                for child in li.children:
                    if isinstance(child, NavigableString):
                        p.add_run(str(child))
                    elif isinstance(child, Tag):
                        r = p.add_run(child.get_text())
                        if child.name in ["strong", "b"]:
                            r.font.bold = True
                        if child.name in ["em", "i"]:
                            r.font.italic = True
                        if child.name == "a":
                            r.font.underline = True
                            r.font.color.rgb = COLOR_AMBER
                for r in p.runs:
                    r.font.name = "Calibri"
                    r.font.size = Pt(10)
                    if not r.font.color.rgb:
                        r.font.color.rgb = COLOR_DARK

        elif tag_name == "table":
            render_table_element(doc, elem)

        elif tag_name == "hr":
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            r = p.add_run("―" * 45)
            r.font.color.rgb = RGBColor(203, 213, 225)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER

def build_word_document(frontmatter, body, lang="vi"):
    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run("THE RICE TOUR — BẾN THÀNH TOPIC CLUSTER 2026" if lang == "vi" else "THE RICE TOUR — BEN THANH HERITAGE COMPENDIUM 2026")
        hrun.font.name = "Calibri"
        hrun.font.size = Pt(8.5)
        hrun.font.color.rgb = COLOR_SLATE

        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        frun = fp.add_run("thericetour.com | Du Lịch Có GUU" if lang == "vi" else "thericetour.com | Bespoke Inbound Journeys")
        frun.font.name = "Calibri"
        frun.font.size = Pt(8.5)
        frun.font.color.rgb = COLOR_SLATE

    title = frontmatter.get("title", "Cẩm Nang Du Lịch Bến Thành")
    tp = doc.add_paragraph()
    tp.paragraph_format.space_before = Pt(14)
    tp.paragraph_format.space_after = Pt(6)
    trun = tp.add_run(title)
    trun.font.name = "Calibri"
    trun.font.size = Pt(21)
    trun.font.bold = True
    trun.font.color.rgb = COLOR_NAVY

    meta_table = doc.add_table(rows=1, cols=1)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_cell = meta_table.cell(0, 0)
    meta_cell.width = Inches(6.5)
    set_cell_background(meta_cell, "F8FAFC")
    set_cell_margins(meta_cell, top=100, bottom=100, left=160, right=160)

    tcPr = meta_cell._element.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="none"/>'
        f'<w:bottom w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'<w:left w:val="single" w:sz="18" w:space="0" w:color="{NAVY_HEX}"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)

    mp = meta_cell.paragraphs[0]
    mp.paragraph_format.space_before = Pt(0)
    mp.paragraph_format.space_after = Pt(0)
    mp.paragraph_format.line_spacing = 1.2
    
    cats = frontmatter.get("categories", "Cẩm Nang Sài Gòn")
    tags = frontmatter.get("tags", "Bến Thành, Quận 1")
    pub_date = frontmatter.get("published_date", "2026-09-07")
    author = frontmatter.get("author", "The Rice Tour Editorial")
    read_time = frontmatter.get("read_time", "12")

    meta_text = (
        f"🏷️ {'Chuyên mục' if lang=='vi' else 'Categories'}: {cats}  |  ⏱️ {'Thời gian đọc' if lang=='vi' else 'Read time'}: {read_time} {'phút' if lang=='vi' else 'mins'}\n"
        f"📅 {'Cập nhật' if lang=='vi' else 'Published'}: {pub_date[:10]}  |  ✍️ {'Tác giả' if lang=='vi' else 'Editorial'}: {author}\n"
        f"🏷️ {'Thẻ từ khóa' if lang=='vi' else 'Tags'}: {tags}"
    )
    mrun = mp.add_run(meta_text)
    mrun.font.name = "Calibri"
    mrun.font.size = Pt(9)
    mrun.font.color.rgb = COLOR_SLATE

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # Convert markdown body to HTML with tables & extras
    html_body = markdown.markdown(body, extensions=["tables", "fenced_code", "nl2br"])
    parse_html_to_docx(doc, html_body)

    return doc

def main():
    vi_input_dir = "/Users/huynhtronghieu/Documents/thericetour/content-pipeline/campaign-ben-thanh/vietnamese"
    en_input_dir = "/Users/huynhtronghieu/Documents/thericetour/content-pipeline/campaign-ben-thanh/english"
    
    out_campaign_word = "/Users/huynhtronghieu/Documents/thericetour/content-pipeline/campaign-ben-thanh/word"
    out_general_vi = "/Users/huynhtronghieu/Documents/thericetour/content-pipeline/06-word-documents/01-vietnamese-word"
    out_general_en = "/Users/huynhtronghieu/Documents/thericetour/content-pipeline/06-word-documents/02-english-word"
    out_general_base = "/Users/huynhtronghieu/Documents/thericetour/content-pipeline/06-word-documents"

    os.makedirs(out_campaign_word, exist_ok=True)
    os.makedirs(out_general_vi, exist_ok=True)
    os.makedirs(out_general_en, exist_ok=True)

    all_18_vi_files = sorted([f for f in os.listdir(vi_input_dir) if f.endswith(".md") and f.startswith("0")])
    all_18_en_files = sorted([f for f in os.listdir(en_input_dir) if f.endswith(".md") and f.startswith("0")])

    print(f"\n=======================================================")
    print(f" XUẤT BẢN WORD (.docx) TOÀN BỘ 18 BÀI VIẾT BẾN THÀNH")
    print(f"=======================================================")

    # 1. Export 18 Vietnamese individual docs
    print(f"\n[PHASE 1] Xuất bản {len(all_18_vi_files)} bài viết Tiếng Việt sang Word...")
    for idx, f in enumerate(all_18_vi_files):
        in_path = os.path.join(vi_input_dir, f)
        fm, body = parse_markdown_file(in_path)
        doc = build_word_document(fm, body, lang="vi")
        out_name = f.replace('.md', '.docx')
        
        # Save to campaign folder
        doc.save(os.path.join(out_campaign_word, out_name))
        # Save to general pipeline folder
        doc.save(os.path.join(out_general_vi, out_name))
        print(f"  [VI {idx+1}/{len(all_18_vi_files)}] Đã xuất: {out_name}")

    # 2. Master VI Consolidated Dossier
    print("\n[PHASE 2] Tạo File Master Dossier Tổng Hợp 18 Bài Tiếng Việt...")
    master_vi = Document()
    for s in master_vi.sections:
        s.top_margin = Inches(0.8)
        s.bottom_margin = Inches(0.8)
        s.left_margin = Inches(0.8)
        s.right_margin = Inches(0.8)
        s.header.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
        s.header.paragraphs[0].add_run("THE RICE TOUR — TỔNG HỢP 18 BÀI VIẾT BẾN THÀNH 2026").font.size = Pt(8.5)
        s.footer.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        s.footer.paragraphs[0].add_run("thericetour.com | Cẩm Nang Du Lịch Có GUU").font.size = Pt(8.5)

    cp = master_vi.add_paragraph()
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp.paragraph_format.space_before = Pt(120)
    cp.paragraph_format.space_after = Pt(16)
    c1 = cp.add_run("THE RICE TOUR\nBỘ SƯU TẬP TOÀN TẬP 18 CẨM NANG THỰC ĐỊA\nCHIẾN DỊCH TOPIC CLUSTER BẾN THÀNH 2026")
    c1.font.name = "Calibri"
    c1.font.size = Pt(24)
    c1.font.bold = True
    c1.font.color.rgb = COLOR_NAVY

    cp2 = master_vi.add_paragraph()
    cp2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp2.paragraph_format.space_after = Pt(40)
    c2 = cp2.add_run("Topic Cluster Master Dossier — 18 Chuyên Đề Khám Phá Trái Tim Sài Gòn Quận 1\nBao gồm: Di Sản Văn Hóa, Ẩm Thực, Metro Số 1, Đi Bộ, Sky Bar, Khách Sạn Boutique,\nPillar Triad (Có gì chơi - Kinh nghiệm A-Z) và Cụm Giải Quyết Vấn Đề Nóng (Lừa đảo, Đổi ngoại tệ Hà Tâm, Bãi giữ xe, Sân bay Tân Sơn Nhất)")
    c2.font.name = "Calibri"
    c2.font.size = Pt(11)
    c2.font.italic = True
    c2.font.color.rgb = COLOR_SLATE

    master_vi.add_page_break()

    for idx, f in enumerate(all_18_vi_files):
        in_path = os.path.join(vi_input_dir, f)
        fm, body = parse_markdown_file(in_path)
        art_doc = build_word_document(fm, body, lang="vi")
        for elem in art_doc.paragraphs:
            p = master_vi.add_paragraph(elem.text, style=elem.style.name if elem.style else 'Normal')
            p.paragraph_format.space_before = elem.paragraph_format.space_before
            p.paragraph_format.space_after = elem.paragraph_format.space_after
            p.paragraph_format.line_spacing = elem.paragraph_format.line_spacing
            p.alignment = elem.alignment
            for r in elem.runs:
                nr = p.add_run(r.text)
                nr.font.name = r.font.name
                nr.font.size = r.font.size
                nr.font.bold = r.font.bold
                nr.font.italic = r.font.italic
                if r.font.color and r.font.color.rgb:
                    nr.font.color.rgb = r.font.color.rgb
        if idx < len(all_18_vi_files) - 1:
            master_vi.add_page_break()

    master_vi_path1 = os.path.join(out_campaign_word, "000_TONG_HOP_18_BAI_VIET_BEN_THANH_VI.docx")
    master_vi_path2 = os.path.join(out_general_base, "000_TONG_HOP_18_BAI_VIET_BEN_THANH_VI.docx")
    master_vi.save(master_vi_path1)
    master_vi.save(master_vi_path2)
    print(f"  -> Đã lưu Master Dossier Tiếng Việt tại: {master_vi_path1}")

    # 3. Export 18 English individual docs
    print(f"\n[PHASE 3] Xuất bản {len(all_18_en_files)} bài viết Tiếng Anh sang Word...")
    for idx, f in enumerate(all_18_en_files):
        in_path = os.path.join(en_input_dir, f)
        fm, body = parse_markdown_file(in_path)
        doc = build_word_document(fm, body, lang="en")
        out_name = f.replace('.md', '.docx')
        
        doc.save(os.path.join(out_campaign_word, out_name))
        doc.save(os.path.join(out_general_en, out_name))
        print(f"  [EN {idx+1}/{len(all_18_en_files)}] Successfully exported: {out_name}")

    # 4. Master EN Consolidated Dossier
    print("\n[PHASE 4] Creating Master Consolidated English Compendium (18 Guides)...")
    master_en = Document()
    for s in master_en.sections:
        s.top_margin = Inches(0.8)
        s.bottom_margin = Inches(0.8)
        s.left_margin = Inches(0.8)
        s.right_margin = Inches(0.8)
        s.header.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
        s.header.paragraphs[0].add_run("THE RICE TOUR — BEN THANH 18 HERITAGE GUIDES 2026").font.size = Pt(8.5)
        s.footer.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        s.footer.paragraphs[0].add_run("thericetour.com | Bespoke Inbound Journeys").font.size = Pt(8.5)

    ep = master_en.add_paragraph()
    ep.alignment = WD_ALIGN_PARAGRAPH.CENTER
    ep.paragraph_format.space_before = Pt(120)
    ep.paragraph_format.space_after = Pt(16)
    e1 = ep.add_run("THE RICE TOUR\nTHE DEFINITIVE BEN THANH TRAVEL COMPENDIUM\n18 FIELD GUIDES, SURVIVAL MANUALS & LOGISTICS (2026)")
    e1.font.name = "Calibri"
    e1.font.size = Pt(24)
    e1.font.bold = True
    e1.font.color.rgb = COLOR_NAVY

    ep2 = master_en.add_paragraph()
    ep2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    ep2.paragraph_format.space_after = Pt(40)
    e2 = ep2.add_run("A Curated Inbound Compendium of District 1’s Epicenter: Colonial Relics, Culinary Alleyways,\nMetro Line 1 Logistics, Secret Apartment Cafes, Twilight Skybars, Boutique Stays,\nIntent Pillar Guides (Things to Do & A-Z Survival) and Vital Pain Points (Scams, Ha Tam FX, Parking & Airport Transfers)")
    e2.font.name = "Calibri"
    e2.font.size = Pt(11)
    e2.font.italic = True
    e2.font.color.rgb = COLOR_SLATE

    master_en.add_page_break()

    for idx, f in enumerate(all_18_en_files):
        in_path = os.path.join(en_input_dir, f)
        fm, body = parse_markdown_file(in_path)
        art_doc = build_word_document(fm, body, lang="en")
        for elem in art_doc.paragraphs:
            p = master_en.add_paragraph(elem.text, style=elem.style.name if elem.style else 'Normal')
            p.paragraph_format.space_before = elem.paragraph_format.space_before
            p.paragraph_format.space_after = elem.paragraph_format.space_after
            p.paragraph_format.line_spacing = elem.paragraph_format.line_spacing
            p.alignment = elem.alignment
            for r in elem.runs:
                nr = p.add_run(r.text)
                nr.font.name = r.font.name
                nr.font.size = r.font.size
                nr.font.bold = r.font.bold
                nr.font.italic = r.font.italic
                if r.font.color and r.font.color.rgb:
                    nr.font.color.rgb = r.font.color.rgb
        if idx < len(all_18_en_files) - 1:
            master_en.add_page_break()

    master_en_path1 = os.path.join(out_campaign_word, "000_FULL_COLLECTION_18_ARTICLES_BEN_THANH_EN.docx")
    master_en_path2 = os.path.join(out_general_base, "000_FULL_COLLECTION_18_ARTICLES_BEN_THANH_EN.docx")
    master_en.save(master_en_path1)
    master_en.save(master_en_path2)
    print(f"  -> Successfully saved Master English Compendium at: {master_en_path1}")

    print("\n=======================================================")
    print(" HOÀN TẤT XUẤT BẢN TOÀN BỘ 36 FILE WORD LẺ + 2 MASTER DOSSIER")
    print(f" 📁 Chiến dịch riêng biệt: {out_campaign_word}")
    print(f" 📁 Thư viện tổng: {out_general_base}")
    print("=======================================================\n")

if __name__ == "__main__":
    main()
