import os
import re
import openpyxl

wb = openpyxl.load_workbook('content-pipeline/ben-thanh.xlsx')
sheet = wb.active

embeds_by_code = {}
for r in range(2, sheet.max_row + 1):
    code = str(sheet.cell(r, 1).value).strip().zfill(3)
    fname = str(sheet.cell(r, 2).value).strip()
    topic = str(sheet.cell(r, 3).value).strip()
    user = str(sheet.cell(r, 4).value).strip()
    link = str(sheet.cell(r, 5).value).strip()
    embed = str(sheet.cell(r, 7).value or '').strip()
    if code not in embeds_by_code:
        embeds_by_code[code] = []
    embeds_by_code[code].append({
        'topic': topic,
        'user': user,
        'link': link,
        'embed': embed
    })

featured_images = {
    "001": "https://media.thericetour.com/uploads/ben-thanh-market-clock-tower.webp",
    "002": "https://media.thericetour.com/uploads/ho-chi-minh-city-museum-of-fine-arts.webp",
    "003": "https://media.thericetour.com/uploads/ben-thanh-market-street-food.webp",
    "004": "https://media.thericetour.com/uploads/central-ho-chi-minh-city-street-scene.webp",
    "005": "https://media.thericetour.com/uploads/reunification-palace-saigon.webp",
    "006": "https://media.thericetour.com/uploads/ben-thanh-metro-station-circular-entrance.webp",
    "007": "https://media.thericetour.com/uploads/mariamman-hindu-temple-saigon.webp",
    "008": "https://media.thericetour.com/uploads/ben-thanh-market-shopping.webp",
    "009": "https://media.thericetour.com/uploads/saigon-hop-on-hop-off-bus.webp",
    "010": "https://media.thericetour.com/uploads/apartment-cafe.webp",
    "011": "https://media.thericetour.com/uploads/rooftop-bar-ben-thanh-market-view.webp",
    "012": "https://media.thericetour.com/uploads/hotel-continental-saigon.webp",
    "013": "https://media.thericetour.com/uploads/ben-thanh-market-atmosphere.webp",
    "014": "https://media.thericetour.com/uploads/ben-thanh-market-main-gate.webp",
    "015": "https://media.thericetour.com/uploads/ben-thanh-market-tourist-tips.webp",
    "016": "https://media.thericetour.com/uploads/currency-exchange-near-ben-thanh-market-1.webp",
    "017": "https://media.thericetour.com/uploads/ben-thanh-market-motorbike-parking.webp",
    "018": "https://media.thericetour.com/uploads/tan-son-nhat-airport.webp"
}

def wrap_embed(embed_html):
    clean_embed = embed_html.strip()
    return f'\n<div class="instagram-embed-wrapper my-8 flex justify-center not-prose w-full">\n  <div class="w-full max-w-[540px] overflow-hidden rounded-2xl shadow-sm border border-slate-200/80 bg-white p-2">\n{clean_embed}\n  </div>\n</div>\n'

def update_frontmatter_image(content, img_url):
    if not content.startswith('---'):
        return content
    parts = content.split('---', 2)
    if len(parts) < 3:
        return content
    fm = parts[1]
    body = parts[2]
    
    # Check if featured_image already exists
    if re.search(r'^featured_image:\s*.*$', fm, re.MULTILINE):
        fm = re.sub(r'^featured_image:\s*.*$', f'featured_image: "{img_url}"', fm, flags=re.MULTILINE)
    else:
        fm += f'\nfeatured_image: "{img_url}"\n'
    return f'---{fm}---{body}'

def process_file(filepath, is_english=True):
    basename = os.path.basename(filepath)
    code = basename[:3]
    if code not in embeds_by_code:
        return
    
    with open(filepath, 'r', encoding='utf-8') as f:
        content = f.read()
        
    img_url = featured_images.get(code)
    if img_url:
        content = update_frontmatter_image(content, img_url)
        
    embed_items = embeds_by_code[code]
    
    # Avoid duplicate insertions
    if 'instagram-embed-wrapper' in content:
        # Already has embeds, skip body insertion
        with open(filepath, 'w', encoding='utf-8') as f:
            f.write(content)
        return
        
    if code == '001':
        # 5 embeds into 5 subsections
        # In EN: 2.1 Ben Thanh, 2.2 Museum of Fine Arts, 2.3 Independence Palace, 2.4 Mariamman, 2.5 Metro
        # In VI: 2.1 Chợ Bến Thành, 2.2 Bảo tàng Mỹ thuật, 2.3 Dinh Độc Lập, 2.4 Đền Hindu, 2.5 Ga Metro
        subsections = [
            (r'(###\s*2\.1\..*?\n\n.*?\n\n)', 0),
            (r'(###\s*2\.2\..*?\n\n.*?\n\n)', 1),
            (r'(###\s*2\.3\..*?\n\n.*?\n\n)', 2),
            (r'(###\s*2\.4\..*?\n\n.*?\n\n)', 3),
            (r'(###\s*2\.5\..*?\n\n.*?\n\n)', 4),
        ]
        for pattern, embed_idx in subsections:
            match = re.search(pattern, content)
            if match and embed_idx < len(embed_items):
                embed_html = wrap_embed(embed_items[embed_idx]['embed'])
                pos = match.end()
                content = content[:pos] + embed_html + content[pos:]
    else:
        # Single embed into section 2
        # Match '## 2. ...' and insert after first paragraph
        match = re.search(r'(##\s*2\..*?\n\n.*?\n\n)', content)
        if not match:
            # Fallback to after ## 3.
            match = re.search(r'(##\s*3\..*?\n\n.*?\n\n)', content)
        if not match:
            # Fallback to after first ##
            match = re.search(r'(##\s*.*?\n\n.*?\n\n)', content)
            
        if match:
            embed_html = wrap_embed(embed_items[0]['embed'])
            pos = match.end()
            content = content[:pos] + embed_html + content[pos:]
            
    with open(filepath, 'w', encoding='utf-8') as f:
        f.write(content)
    print(f'Updated: {filepath}')

# Apply to all relevant directories
dirs = [
    ('content-pipeline/04-english', True),
    ('content-pipeline/campaign-ben-thanh/english', True),
    ('content-pipeline/02-vietnamese-guu', False),
    ('content-pipeline/03-qa-passed', False),
    ('content-pipeline/campaign-ben-thanh/vietnamese', False)
]

for d, is_en in dirs:
    if not os.path.exists(d):
        continue
    for fname in sorted(os.listdir(d)):
        if fname.endswith('.md') and fname[:3].isdigit() and int(fname[:3]) <= 18:
            process_file(os.path.join(d, fname), is_en)

print("All markdown files updated successfully!")
